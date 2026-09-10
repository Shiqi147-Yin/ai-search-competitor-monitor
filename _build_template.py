"""生成 Excel 模板文件（运行一次即可）"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUTPUT_PATH = Path(__file__).parent / "templates" / "competitor_updates_template.xlsx"

HEADERS = ["发布时间", "竞品", "动态标题", "动态概述", "原始链接", "信息平台", "备注"]
EXAMPLE_ROW = [
    "2026-07-20",
    "Tavily",
    "Tavily 发布新版 Search API，支持实时网页内容抓取",
    "新版 API 提供 include_raw_content 参数，可直接返回完整网页正文，适合 RAG 场景。",
    "https://docs.tavily.com/changelog/2026-07-20",
    "官网",
    "Querit 尚不支持原始内容返回，需评估是否跟进",
]

COMPETITOR_OPTIONS = "Tavily,Exa,Brave,Querit,Other"
PLATFORM_OPTIONS = "官网,Blog,GitHub,X,LinkedIn,Discord,Event,Other"

COL_WIDTHS = [14, 12, 40, 50, 45, 12, 40]

HEADER_FILL = PatternFill("solid", fgColor="2C5F8A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
EXAMPLE_FILL = PatternFill("solid", fgColor="EBF3FA")
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


def build_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竞品动态"

    # 写表头
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    # 写示例行
    for col_idx, value in enumerate(EXAMPLE_ROW, start=1):
        cell = ws.cell(row=2, column=col_idx, value=value)
        cell.fill = EXAMPLE_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = THIN_BORDER

    # 调整列宽和行高
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 60
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    # 下拉验证：竞品（B列）
    dv_competitor = DataValidation(
        type="list",
        formula1=f'"{COMPETITOR_OPTIONS}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="输入错误",
        error="请从下拉列表中选择竞品名称",
    )
    ws.add_data_validation(dv_competitor)
    dv_competitor.sqref = "B2:B1000"

    # 下拉验证：信息平台（F列）
    dv_platform = DataValidation(
        type="list",
        formula1=f'"{PLATFORM_OPTIONS}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="输入错误",
        error="请从下拉列表中选择信息平台",
    )
    ws.add_data_validation(dv_platform)
    dv_platform.sqref = "F2:F1000"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"模板已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    build_template()
